"""Replace the placeholder plagiarism chapter with the Plagiarism.Free result table."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = r"C:\Users\User\Desktop\research project\PROJECT_REPORT.docx"
URL = "https://plagiarism.free/check/6999526b5d384f7d95a525f888b5175d/"


def set_run(run, size=12, bold=False, white=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0)


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tb.append(el)
    tcPr.append(tb)


def main():
    doc = Document(SRC)
    heading = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("15. PLAGIARISM"):
            heading = p
            break
    if heading is None:
        raise SystemExit("heading not found")

    el = heading._element.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt

    def add_para(text, size=12, bold=False, align="justify", before=6, after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = {
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
        }[align]
        if text:
            r = p.add_run(text)
            set_run(r, size=size, bold=bold)
        return p

    bits = []
    bits.append(
        add_para(
            "The project report was checked with Plagiarism.Free (Web and Academic Index). "
            "The official result for this manuscript is recorded in Table 15.1."
        )._element
    )
    bits.append(
        add_para(
            "We hereby declare that the contents of this research project report are our original work "
            "except where citations are provided. The similarity index obtained from Plagiarism.Free on "
            "7 October 2026 is 1%, which is within the limit prescribed by the MCA Department, "
            "P.E.S.'s Modern College of Engineering, Pune. The full report is available at " + URL
        )._element
    )
    bits.append(
        add_para(
            "Table 15.1  Plagiarism verification result (Plagiarism.Free)",
            size=12,
            bold=True,
            align="center",
            before=12,
            after=4,
        )._element
    )

    rows = [
        ["Verification Parameter", "Institutional Limit / Norm", "Plagiarism.Free Official Result"],
        ["Verification Engine", "Standard Academic Checker", "Plagiarism.Free (Web & Academic Index)"],
        ["Verification Timestamp", "Official Record", "2026-10-07 15:37 UTC"],
        ["Manuscript Words Examined", "Full Research Paper Text", "8,743 words evaluated"],
        ["Overall Similarity Index (Matched)", "Maximum permissible: \u226415%", "1% matched (99% Unique)"],
        ["Internet & Live Web Matches", "Maximum permissible: \u226410%", "Included in the 1% overall index (4 sources listed)"],
        ["Published Academic Papers / Theses", "Maximum permissible: \u22648%", "Included in the 1% overall index"],
        ["Originality Assessment Rating", "Authentic Manuscript", "Looks original (99% Unique)"],
        ["Plagiarism Clearance Status", "Within the department limit", "Within the permissible limit (\u226415%)"],
        ["Report link", "Public verification page", URL],
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.line_spacing = 1.0
            if i == 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(val)
            set_run(r, size=11, bold=(i == 0 or j == 0), white=(i == 0))
            if i == 0:
                shade(cell, "1F2933")
            borders(cell)
    bits.append(table._tbl)

    cursor = heading._element
    for item in bits:
        cursor.addnext(item)
        cursor = item

    doc.save(SRC)
    print("saved", SRC)


if __name__ == "__main__":
    main()
