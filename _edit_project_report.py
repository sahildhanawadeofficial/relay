"""Edit PROJECT_REPORT.docx: diagrams, TOC page numbers, certificate."""
from __future__ import annotations

import copy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor, Twips

ROOT = Path(r"C:\Users\User\Desktop\research project")
SRC = Path(r"C:\Users\User\Downloads\PROJECT_REPORT.docx")
DST = SRC  # update the attached file
COPY = ROOT / "PROJECT_REPORT.docx"
FIG = ROOT / "report-diagrams"
LOGO = ROOT / "college-logo.jpg"
FONT = r"C:\Windows\Fonts\times.ttf"
FONTB = r"C:\Windows\Fonts\timesbd.ttf"
FONTI = r"C:\Windows\Fonts\timesi.ttf"
DPI = 300


def F(path, size):
    return ImageFont.truetype(path, size)


def text_size(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def wrap(draw, text, font, max_w):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = w if not cur else cur + " " + w
        if text_size(draw, trial, font)[0] <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


class Canvas:
    def __init__(self, w, h):
        self.im = Image.new("RGB", (w, h), "white")
        self.d = ImageDraw.Draw(self.im)
        self.w, self.h = w, h

    def save(self, name):
        FIG.mkdir(exist_ok=True)
        path = FIG / name
        self.im.save(path, "PNG", dpi=(DPI, DPI))
        return path

    def round_box(self, xy, fill="#F4F6F8", outline="#1F2933", radius=16, width=3):
        self.d.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)

    def center_lines(self, xy, lines, font, fill="#1F2933", gap=6):
        x0, y0, x1, y1 = xy
        heights = [text_size(self.d, ln, font)[1] for ln in lines]
        total = sum(heights) + gap * (len(lines) - 1)
        y = (y0 + y1 - total) / 2
        for ln, th in zip(lines, heights):
            tw, _ = text_size(self.d, ln, font)
            self.d.text(((x0 + x1 - tw) / 2, y), ln, font=font, fill=fill)
            y += th + gap

    def arrow(self, a, b, color="#1F2933", width=3):
        self.d.line([a, b], fill=color, width=width)
        x0, y0 = a
        x1, y1 = b
        import math
        ang = math.atan2(y1 - y0, x1 - x0)
        L, W = 16, 8
        p1 = (x1, y1)
        p2 = (x1 - L * math.cos(ang) + W * math.sin(ang), y1 - L * math.sin(ang) - W * math.cos(ang))
        p3 = (x1 - L * math.cos(ang) - W * math.sin(ang), y1 - L * math.sin(ang) + W * math.cos(ang))
        self.d.polygon([p1, p2, p3], fill=color)

    def label(self, xy, text, font, fill="#1F2933"):
        self.d.text(xy, text, font=font, fill=fill)


def fig_layers():
    c = Canvas(2100, 1500)
    title = F(FONTB, 48)
    body = F(FONT, 36)
    c.label((80, 40), "Logical layers of Relay", title)
    layers = [
        ("1  Presentation", "Landing, login, dashboard, chatbot workspace, Shadow DOM widget"),
        ("2  API", "Session routes  /api/chatbots/*     Public routes  /api/public/*"),
        ("3  Domain", "answerChatbotQuery, parser, chunker, API key, CORS"),
        ("4  Data", "MongoDB: users and chatbot config     Pinecone: vectors and metadata"),
        ("5  Inference", "Pinecone hosted embeddings     OpenRouter chat completions"),
    ]
    y = 140
    for i, (h, sub) in enumerate(layers):
        fill = "#E8EEF6" if i % 2 == 0 else "#F7F1E8"
        c.round_box((160, y, 1940, y + 220), fill=fill, radius=18)
        c.center_lines((160, y, 1940, y + 220), [h, sub], body if False else F(FONTB, 40))
        # overwrite with two fonts
        c.d.rectangle((161, y + 1, 1939, y + 219), fill=fill)
        c.round_box((160, y, 1940, y + 220), fill=fill, radius=18)
        c.center_lines((180, y + 20, 1920, y + 120), [h], F(FONTB, 42))
        c.center_lines((180, y + 110, 1920, y + 200), [sub], F(FONT, 32), fill="#333333")
        if i < len(layers) - 1:
            c.arrow((1050, y + 220), (1050, y + 268))
        y += 270
    return c.save("fig_5_1_layers.png")


def fig_ingestion():
    c = Canvas(1400, 2100)
    body = F(FONTB, 36)
    steps = [
        "Owner selects PDF / DOCX / TXT",
        "Client slices file at 3.5 MB",
        "POST /api/chatbots/{id}/documents",
        "Session and ownership check",
        "Extract text",
        "Recursive chunking\n800 chars, overlap 150",
        "Pinecone embed\ninputType = passage",
        "Upsert vectors with chatbot_id",
        "Return chunk count to the UI",
    ]
    y = 40
    for i, s in enumerate(steps):
        lines = s.split("\n")
        c.round_box((120, y, 1280, y + 160), fill="#E7F2EA", radius=28)
        c.center_lines((140, y, 1260, y + 160), lines, body)
        if i < len(steps) - 1:
            c.arrow((700, y + 160), (700, y + 200))
        y += 220
    return c.save("fig_5_2_ingestion.png")


def fig_query():
    c = Canvas(2100, 1680)
    b = F(FONTB, 34)
    sm = F(FONT, 30)
    def box(xy, lines, fill):
        c.round_box(xy, fill=fill, radius=20)
        c.center_lines(xy, lines, b)
    box((700, 30, 1400, 150), ["Question received"], "#E8EEF6")
    c.arrow((1050, 150), (1050, 190))
    box((620, 190, 1480, 320), ["Authenticate", "session or API key + origin"], "#E8EEF6")
    c.arrow((1050, 320), (1050, 360))
    box((620, 360, 1480, 490), ["Embed query", "inputType = query"], "#F7F1E8")
    c.arrow((1050, 490), (1050, 530))
    box((560, 530, 1540, 680), ["Pinecone top-k", "filter chatbot_id"], "#F7F1E8")
    c.arrow((1050, 680), (1050, 740))
    # diamond
    cx, cy = 1050, 860
    c.d.polygon([(cx, cy - 110), (cx + 280, cy), (cx, cy + 110), (cx - 280, cy)], fill="#FFF6D8", outline="#1F2933")
    c.center_lines((cx - 220, cy - 50, cx + 220, cy + 50), ["Matches", "empty?"], b)
    c.arrow((770, 860), (430, 860))
    c.label((560, 800), "Yes", b)
    box((40, 780, 400, 940), ["Return", "\"could not find\""], "#F8E8E8")
    c.arrow((1330, 860), (1680, 860))
    c.label((1360, 800), "No", b)
    box((1680, 760, 2060, 960), ["Build context", "Call OpenRouter"], "#E7F2EA")
    c.arrow((1870, 960), (1870, 1040))
    cx2, cy2 = 1870, 1160
    c.d.polygon([(cx2, cy2 - 100), (cx2 + 230, cy2), (cx2, cy2 + 100), (cx2 - 230, cy2)], fill="#FFF6D8", outline="#1F2933")
    c.center_lines((cx2 - 180, cy2 - 40, cx2 + 180, cy2 + 40), ["Valid", "answer?"], b)
    c.arrow((1640, 1160), (1280, 1160))
    c.label((1400, 1100), "No", b)
    box((760, 1060, 1280, 1260), ["Return quoted", "top chunks"], "#F7F1E8")
    c.arrow((1870, 1260), (1870, 1340))
    c.label((1900, 1280), "Yes", b)
    box((1560, 1340, 2060, 1520), ["Return answer", "and sources"], "#E7F2EA")
    return c.save("fig_5_3_query.png")


def fig_architecture():
    c = Canvas(2200, 1500)
    b = F(FONTB, 34)
    s = F(FONT, 28)
    def box(xy, lines, fill):
        c.round_box(xy, fill=fill, radius=16)
        c.center_lines(xy, lines, b)
    box((80, 40, 1000, 200), ["Chatbot owner", "Dashboard  ·  JWT session"], "#E8EEF6")
    box((1200, 40, 2120, 200), ["Website visitor", "Widget  ·  pk_live_ API key"], "#E8EEF6")
    c.arrow((540, 200), (540, 280))
    c.arrow((1660, 200), (1660, 280))
    box((160, 280, 2040, 620), ["Next.js 15 application (Relay)", "Auth.js   ·   parser   ·   chunker   ·   answerChatbotQuery", "Session APIs /api/chatbots/*     Public APIs /api/public/*"], "#EEF2F6")
    c.arrow((500, 620), (500, 720))
    c.arrow((1100, 620), (1100, 720))
    c.arrow((1700, 620), (1700, 720))
    box((80, 720, 680, 980), ["MongoDB", "Users, chatbots", "API keys, widget config"], "#F7F1E8")
    box((760, 720, 1440, 1080), ["Pinecone", "Serverless index", "Hosted E5 embeddings", "Filter on chatbot_id"], "#E7F2EA")
    box((1520, 720, 2120, 980), ["OpenRouter", "LLM completion", "Model fallbacks"], "#F8E8E8")
    c.label((80, 1180), "Both the dashboard and the widget call the same retrieval function.", s)
    return c.save("fig_6_1_architecture.png")


def stick(c, x, y, label, font):
    # head, body, arms, legs
    c.d.ellipse((x - 28, y, x + 28, y + 56), outline="#1F2933", width=3)
    c.d.line((x, y + 56, x, y + 150), fill="#1F2933", width=3)
    c.d.line((x - 48, y + 90, x + 48, y + 90), fill="#1F2933", width=3)
    c.d.line((x, y + 150, x - 40, y + 220), fill="#1F2933", width=3)
    c.d.line((x, y + 150, x + 40, y + 220), fill="#1F2933", width=3)
    tw, th = text_size(c.d, label, font)
    c.d.text((x - tw / 2, y + 240), label, font=font, fill="#1F2933")


def fig_usecase():
    c = Canvas(2400, 1700)
    title = F(FONTB, 36)
    sm = F(FONT, 30)
    stick(c, 150, 360, "Chatbot Owner", title)
    stick(c, 2250, 360, "Website Visitor", title)
    c.d.rectangle((420, 70, 1980, 1580), outline="#1F2933", width=4)
    tw, _ = text_size(c.d, "Relay", F(FONTB, 42))
    c.d.text(((420 + 1980 - tw) / 2, 90), "Relay", font=F(FONTB, 42), fill="#1F2933")

    def oval(x, y, w, h, lab):
        c.d.ellipse((x, y, x + w, y + h), outline="#1F2933", width=3, fill="#F4F7FB")
        c.center_lines((x + 20, y, x + w - 20, y + h), wrap(c.d, lab, sm, w - 50), sm)

    owner = [
        "Sign in with Google",
        "Create, list, delete chatbot",
        "Upload document",
        "Ask question in dashboard",
        "Rotate API key",
        "Set origins and appearance",
    ]
    y = 180
    for lab in owner:
        oval(500, y, 520, 130, lab)
        c.d.line([(220, 480), (420, 480)], fill="#1F2933", width=3)
        c.d.line([(420, y + 65), (500, y + 65)], fill="#1F2933", width=2)
        y += 155
    visitor = ["Open widget", "Read welcome message", "Ask question and read sources"]
    y = 280
    for lab in visitor:
        oval(1360, y, 520, 130, lab)
        c.d.line([(1880, y + 65), (1980, y + 65)], fill="#1F2933", width=2)
        y += 200
    c.d.line([(2180, 500), (1980, 500)], fill="#1F2933", width=3)
    oval(860, 1360, 520, 140, "Authenticate")
    c.d.line([(760, 1000), (1000, 1360)], fill="#64748B", width=2)
    c.d.line([(1620, 820), (1200, 1360)], fill="#64748B", width=2)
    c.d.text((1080, 1280), "<<include>>", font=sm, fill="#64748B")
    return c.save("fig_6_2_usecase.png")


def fig_sequence(name, actors, messages):
    n = len(messages)
    w, h = 2200, 220 + n * 110
    c = Canvas(w, h)
    head = F(FONTB, 32)
    msgf = F(FONT, 26)
    xs = []
    span = 1900
    left = 150
    for i, a in enumerate(actors):
        x = left + int(i * span / max(1, len(actors) - 1))
        xs.append(x)
        tw, _ = text_size(c.d, a, head)
        c.round_box((x - tw // 2 - 20, 30, x + tw // 2 + 20, 110), fill="#E8EEF6", radius=10)
        c.d.text((x - tw / 2, 52), a, font=head, fill="#1F2933")
        c.d.line((x, 110, x, h - 40), fill="#94A3B8", width=2)
    y = 160
    for src, dst, label in messages:
        x1, x2 = xs[src], xs[dst]
        c.arrow((x1, y), (x2, y))
        mx = (x1 + x2) / 2
        tw, th = text_size(c.d, label, msgf)
        c.d.rectangle((mx - tw / 2 - 8, y - th - 14, mx + tw / 2 + 8, y - 6), fill="white")
        c.d.text((mx - tw / 2, y - th - 12), label, font=msgf, fill="#1F2933")
        y += 110
    return c.save(name)


def fig_erd():
    c = Canvas(2200, 1280)
    title = F(FONTB, 36)
    attr = F(FONT, 28)
    small = F(FONTB, 30)

    def entity(x, y, w, name, fields, pk_n=1):
        c.d.rectangle((x, y, x + w, y + 70), fill="#1F2933")
        tw, th = text_size(c.d, name, title)
        c.d.text((x + (w - tw) / 2, y + 16), name, font=title, fill="white")
        top = y + 70
        row_h = 46
        c.d.rectangle((x, top, x + w, top + row_h * len(fields)), outline="#1F2933", width=3)
        for i, f in enumerate(fields):
            yy = top + i * row_h
            c.d.line((x, yy, x + w, yy), fill="#1F2933", width=1)
            prefix = "PK  " if i < pk_n else ("FK  " if f.startswith("userId") or f.startswith("chatbot_id") else "     ")
            c.d.text((x + 16, yy + 8), prefix + f, font=attr, fill="#1F2933")
        return top + row_h * len(fields)

    entity(40, 160, 560, "USER", ["_id", "name", "email (unique)", "passwordHash", "googleId", "image", "createdAt", "updatedAt"])
    entity(820, 80, 640, "CHATBOT", ["_id", "uuid (unique)", "name", "userId", "apiKey (unique)", "allowedOrigins[]", "widgetConfig", "createdAt, updatedAt"])
    entity(1640, 160, 520, "VECTOR RECORD", ["id", "values[1024]", "chatbot_id", "document_id", "document_name", "chunk_id", "text"], pk_n=1)
    c.d.line([(600, 460), (820, 460)], fill="#1F2933", width=4)
    c.d.text((610, 400), "1", font=small, fill="#1F2933")
    c.d.text((660, 400), "owns", font=attr, fill="#1F2933")
    c.d.text((760, 480), "N", font=small, fill="#1F2933")
    c.d.line([(1460, 500), (1640, 500)], fill="#1F2933", width=4)
    c.d.text((1470, 440), "1", font=small, fill="#1F2933")
    c.d.text((1580, 520), "N", font=small, fill="#1F2933")
    c.label((60, 1100), "Vector Record is stored in Pinecone. Cardinality is enforced by the chatbot_id filter, not by a SQL foreign key.", F(FONT, 30))
    return c.save("fig_6_6_erd.png")


def fig_class():
    c = Canvas(2200, 1500)
    title = F(FONTB, 32)
    body = F(FONT, 26)

    def cls(x, y, w, name, fields, methods):
        head_h = 56
        c.d.rectangle((x, y, x + w, y + head_h), fill="#1F2933")
        tw, _ = text_size(c.d, name, title)
        c.d.text((x + (w - tw) / 2, y + 12), name, font=title, fill="white")
        fh = 36 * max(1, len(fields))
        c.d.rectangle((x, y + head_h, x + w, y + head_h + fh), outline="#1F2933", width=2, fill="#F8FAFC")
        for i, f in enumerate(fields):
            c.d.text((x + 12, y + head_h + 6 + i * 36), f, font=body, fill="#1F2933")
        mh = 36 * max(1, len(methods))
        y2 = y + head_h + fh
        c.d.rectangle((x, y2, x + w, y2 + mh), outline="#1F2933", width=2, fill="#FFFFFF")
        for i, m in enumerate(methods):
            c.d.text((x + 12, y2 + 6 + i * 36), m, font=body, fill="#1F2933")

    cls(40, 40, 680, "User", ["+ name: string", "+ email: string", "+ googleId: string"], ["+ validate()"])
    cls(760, 40, 680, "Chatbot", ["+ uuid: string", "+ userId: ObjectId", "+ apiKey: string", "+ allowedOrigins: string[]"], ["+ save()"])
    cls(1480, 40, 680, "ChatWidget", ["+ apiKey: string", "+ messages", "- shadow: ShadowRoot"], ["+ mount()", "+ send()", "+ destroy()"])
    cls(40, 520, 700, "answerChatbotQuery", ["+ chatbotId: string", "+ query: string", "+ topK: int"], ["+ embedQuery()", "+ queryVectors()", "+ generateAnswer()"])
    cls(800, 520, 640, "Document pipeline", ["+ buffer", "+ filename"], ["+ extractText()", "+ chunkDocument()"])
    cls(1500, 520, 650, "Pinecone gateway", ["+ index", "+ namespace"], ["+ embedTexts()", "+ upsertVectors()", "+ queryVectors()"])
    cls(400, 1040, 700, "Public auth", ["+ origin", "+ apiKey"], ["+ authenticatePublicRequest()", "+ generateApiKey()"])
    cls(1200, 1040, 700, "LLM", ["+ context", "+ sources"], ["+ generateAnswer()"])
    # associations
    c.arrow((720, 180), (760, 180))
    c.d.text((680, 140), "1..*", font=body, fill="#1F2933")
    return c.save("fig_6_7_class.png")


def fig_activity():
    c = Canvas(1600, 2100)
    b = F(FONTB, 32)

    def action(y, text, fill="#E8EEF6"):
        c.round_box((380, y, 1220, y + 110), fill=fill, radius=16)
        c.center_lines((380, y, 1220, y + 110), [text], b)

    def diamond(y, text):
        cx, cy = 800, y + 80
        c.d.polygon([(cx, y), (cx + 280, cy), (cx, y + 160), (cx - 280, cy)], fill="#FFF6D8", outline="#1F2933")
        c.center_lines((cx - 200, cy - 40, cx + 200, cy + 40), [text], b)
        return cy

    # start
    c.d.ellipse((760, 30, 840, 110), fill="#1F2933")
    c.arrow((800, 110), (800, 150))
    action(150, "Authenticate caller")
    c.arrow((800, 260), (800, 300))
    action(300, "Embed the question")
    c.arrow((800, 410), (800, 450))
    action(450, "Query Pinecone with chatbot_id filter", "#F7F1E8")
    c.arrow((800, 560), (800, 610))
    diamond(610, "Matches empty?")
    c.arrow((520, 690), (250, 690))
    c.round_box((40, 630, 250, 760), fill="#F8E8E8", radius=16)
    c.center_lines((40, 630, 250, 760), ["Return", "not found"], b)
    c.d.text((300, 640), "Yes", font=b, fill="#1F2933")
    c.arrow((1080, 690), (1320, 690))
    c.d.text((1100, 640), "No", font=b, fill="#1F2933")
    c.round_box((1320, 620, 1560, 770), fill="#E7F2EA", radius=12)
    c.center_lines((1320, 620, 1560, 770), ["Build", "context"], b)
    c.arrow((1440, 770), (1440, 840))
    c.round_box((1180, 840, 1560, 960), fill="#E7F2EA", radius=12)
    c.center_lines((1180, 840, 1560, 960), ["Call OpenRouter"], b)
    c.arrow((1180, 900), (980, 900))
    diamond(820, "Valid answer?")
    # this diamond overlaps - let me not. The layout got messy.
    return c.save("fig_6_8_activity.png")


def fig_activity_clean():
    c = Canvas(1800, 2200)
    b = F(FONTB, 32)

    def act(x, y, w, h, text, fill="#E8EEF6"):
        c.round_box((x, y, x + w, y + h), fill=fill, radius=18)
        lines = text.split("\n")
        c.center_lines((x, y, x + w, y + h), lines, b)

    def dia(cx, cy, text):
        c.d.polygon([(cx, cy - 90), (cx + 250, cy), (cx, cy + 90), (cx - 250, cy)], fill="#FFF6D8", outline="#1F2933")
        c.center_lines((cx - 180, cy - 36, cx + 180, cy + 36), text.split("\n"), b)

    c.d.ellipse((840, 20, 920, 100), fill="#1F2933")
    c.arrow((880, 100), (880, 140))
    act(480, 140, 800, 110, "Authenticate caller")
    c.arrow((880, 250), (880, 290))
    act(480, 290, 800, 110, "Embed the question")
    c.arrow((880, 400), (880, 440))
    act(400, 440, 960, 120, "Query Pinecone\nfilter chatbot_id", "#F7F1E8")
    c.arrow((880, 560), (880, 640))
    dia(880, 730, "Matches\nempty?")
    c.d.text((620, 640), "Yes", font=b, fill="#1F2933")
    c.arrow((630, 730), (360, 730))
    act(40, 670, 320, 130, "Return\n\"could not find\"", "#F8E8E8")
    c.d.ellipse((150, 830, 220, 900), outline="#1F2933", width=4)
    c.d.ellipse((166, 846, 204, 884), fill="#1F2933")
    c.arrow((180, 800), (180, 830))
    c.d.text((1140, 640), "No", font=b, fill="#1F2933")
    c.arrow((1130, 730), (1400, 730))
    act(1400, 660, 360, 140, "Build context\nfrom chunks", "#E7F2EA")
    c.arrow((1580, 800), (1580, 860))
    act(1360, 860, 420, 120, "Call OpenRouter", "#E7F2EA")
    c.arrow((1360, 920), (1130, 1080))
    dia(880, 1180, "Valid\nanswer?")
    c.d.text((1240, 1125), "Yes", font=b, fill="#1F2933")
    c.arrow((1130, 1180), (1400, 1180))
    act(1400, 1110, 360, 140, "Return answer\nand sources", "#E7F2EA")
    c.d.ellipse((1520, 1280, 1590, 1350), outline="#1F2933", width=4)
    c.d.ellipse((1536, 1296, 1574, 1334), fill="#1F2933")
    c.arrow((1560, 1250), (1560, 1280))
    c.d.text((560, 1280), "No", font=b, fill="#1F2933")
    c.arrow((880, 1270), (880, 1360))
    act(560, 1360, 640, 140, "Return quoted\ntop chunks", "#F7F1E8")
    c.d.ellipse((820, 1540, 900, 1620), outline="#1F2933", width=4)
    c.d.ellipse((840, 1560, 880, 1600), fill="#1F2933")
    c.arrow((880, 1500), (880, 1540))
    return c.save("fig_6_8_activity.png")


def fig_dataflow():
    c = Canvas(2100, 1100)
    b = F(FONTB, 36)
    s = F(FONT, 30)
    c.round_box((60, 40, 1000, 980), fill="#F8FAFC", radius=16)
    c.round_box((1100, 40, 2040, 980), fill="#F8FAFC", radius=16)
    c.center_lines((60, 60, 1000, 140), ["Ingestion  (write)"], b)
    c.center_lines((1100, 60, 2040, 140), ["Query  (read)"], b)
    left = ["Document file", "Extracted text", "Chunks + chatbot_id", "Pinecone upsert", "No LLM call"]
    right = ["User question", "Query embedding", "Filtered top-k chunks", "OpenRouter answer", "Nothing written to the index"]
    y = 180
    for a, r in zip(left, right):
        c.round_box((120, y, 940, y + 120), fill="#E7F2EA", radius=14)
        c.center_lines((120, y, 940, y + 120), [a], s)
        c.round_box((1160, y, 1980, y + 120), fill="#E8EEF6", radius=14)
        c.center_lines((1160, y, 1980, y + 120), [r], s)
        if y < 180 + 4 * 150:
            c.arrow((530, y + 120), (530, y + 148))
            c.arrow((1570, y + 120), (1570, y + 148))
        y += 150
    return c.save("fig_6_9_dataflow.png")


def build_figures():
    paths = {}
    paths["5.1"] = fig_layers()
    paths["5.2"] = fig_ingestion()
    paths["5.3"] = fig_query()
    paths["6.1"] = fig_architecture()
    paths["6.2"] = fig_usecase()
    paths["6.3"] = fig_sequence(
        "fig_6_3_seq_upload.png",
        ["Owner", "Browser", "Next.js API", "Pinecone"],
        [
            (0, 1, "Select PDF / DOCX / TXT"),
            (1, 2, "POST documents (session cookie)"),
            (2, 2, "Check owner, extract, chunk"),
            (2, 3, "Embed passages and upsert"),
            (3, 2, "Upsert acknowledgement"),
            (2, 1, "chunks_processed"),
            (1, 0, "Show document in the list"),
        ],
    )
    paths["6.4"] = fig_sequence(
        "fig_6_4_seq_search.png",
        ["Owner", "Dashboard", "Search API", "Pinecone", "OpenRouter"],
        [
            (0, 1, "Submit question"),
            (1, 2, "POST /search {query, top_k}"),
            (2, 2, "Session and ownership check"),
            (2, 3, "Embed query and filtered top-k"),
            (3, 2, "Matching chunks"),
            (2, 4, "Prompt with question and chunks"),
            (4, 2, "Grounded answer"),
            (2, 1, "{answer, sources}"),
        ],
    )
    paths["6.5"] = fig_sequence(
        "fig_6_5_seq_widget.png",
        ["Visitor", "Widget", "Public API", "RAG pipeline"],
        [
            (1, 2, "OPTIONS preflight"),
            (1, 2, "GET /config  (Bearer pk_live_)"),
            (2, 1, "name, colour, welcome message"),
            (0, 1, "Type a question"),
            (1, 2, "POST /chat {query}"),
            (2, 3, "answerChatbotQuery(uuid)"),
            (3, 2, "{answer, sources}"),
            (2, 1, "Append reply in Shadow DOM"),
        ],
    )
    paths["6.6"] = fig_erd()
    paths["6.7"] = fig_class()
    paths["6.8"] = fig_activity_clean()
    paths["6.9"] = fig_dataflow()
    return paths


# ---------------------------------------------------------------------------
# DOCX editing
# ---------------------------------------------------------------------------
TNR = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
_bm = 100


def set_run_font(run, size=12, bold=False, italic=False, name=TNR, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def fmt(p, size=12, align="left", before=0, after=6, line=1.15, center=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    if align == "center" or center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def para_el_text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def find_para(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def delete_el(el):
    parent = el.getparent()
    parent.remove(el)


def elements_until(start_el, stop_pred):
    out = []
    el = start_el.getnext()
    while el is not None and not stop_pred(el):
        out.append(el)
        el = el.getnext()
    return out


def is_heading_text(el, prefix):
    if el.tag != qn("w:p"):
        return False
    return para_el_text(el).startswith(prefix)


def insert_figure_before(doc, target_el, image_path, caption, width=6.15):
    # picture paragraph
    p = doc.add_paragraph()
    fmt(p, align="center", before=8, after=2, line=1.0)
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    # caption
    cap = doc.add_paragraph()
    fmt(cap, align="center", before=2, after=10, line=1.0)
    r = cap.add_run(caption)
    set_run_font(r, size=11, bold=True, italic=True)
    # move both before target (picture first)
    target_el.addprevious(p._element)
    target_el.addprevious(cap._element)
    # order is wrong: addprevious of cap then... we added p then cap with addprevious
    # last addprevious puts cap immediately before target, and p is before cap. Good:
    # p was addprevious first -> p is before target
    # cap addprevious -> cap is before target, p stays before cap. Result: p, cap, target. Good.


def add_bookmark(paragraph, name):
    global _bm
    _bm += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bm))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bm))
    paragraph._element.insert(0, start)
    paragraph._element.append(end)


def clear_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def rebuild_certificate(doc):
    heading = find_para(doc, "CERTIFICATE")
    ack = find_para(doc, "ACKNOWLEDGEMENT")
    # remove everything between certificate heading and acknowledgement
    for el in elements_until(heading._element, lambda e: e is ack._element):
        delete_el(el)
    # rewrite the certificate heading paragraph: keep its page break, set text
    # The heading paragraph already contains a page break + "CERTIFICATE".
    # We will clear runs but keep page break, then build content AFTER this paragraph
    # and change this paragraph into vertical space? Cleaner: make this paragraph the logo row start.
    # Keep page break by leaving the paragraph and clearing text if break is a separate run.
    el = heading._element
    # Remove text runs but keep breaks
    for child in list(el):
        if child.tag == qn("w:r"):
            has_break = child.find(qn("w:br")) is not None
            has_text = child.find(qn("w:t")) is not None
            if has_text and not has_break:
                el.remove(child)
            elif has_text and has_break:
                for t in child.findall(qn("w:t")):
                    child.remove(t)
    # Insert certificate blocks before ACK
    anchor = ack._element

    def add_p(text, size=12, bold=False, italic=False, align="center", before=0, after=0, line=1.15):
        p = doc.add_paragraph()
        fmt(p, size=size, align=align, before=before, after=after, line=line)
        if text:
            r = p.add_run(text)
            set_run_font(r, size=size, bold=bold, italic=italic)
        anchor.addprevious(p._element)
        return p

    # Header table: logo | college name
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells
    left.text = ""
    lp = left.paragraphs[0]
    fmt(lp, align="center", before=0, after=0, line=1.0)
    lr = lp.add_run()
    lr.add_picture(str(LOGO), width=Inches(1.15))
    clear_cell_borders(left)
    right.text = ""
    lines = [
        ("Progressive Education Society's", 12, False, True),
        ("Modern College of Engineering", 16, True, False),
        ("(An Autonomous Institute Affiliated to Savitribai Phule Pune University)", 12, False, False),
        ("MCA Department", 14, True, False),
    ]
    first = True
    for text, size, bold, italic in lines:
        p = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        p.clear()
        fmt(p, align="center", before=0, after=0, line=1.0)
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)
    clear_cell_borders(right)
    anchor.addprevious(table._tbl)

    add_p("", size=12, before=0, after=0, line=1.0)  # line space = 1
    add_p("CERTIFICATE", size=16, bold=True, before=12, after=0)
    for _ in range(4):  # line space = 4
        add_p("", size=12, line=1.0, after=0, before=0)

    body = doc.add_paragraph()
    fmt(body, align="justify", before=6, after=6, line=1.5)
    r = body.add_run("This is to certify that ")
    set_run_font(r, size=12, bold=False)
    r = body.add_run("AADITI CHORAMLE, CHITRANG CHOUDHARI, PRATIK KAMBLE and SAHIL DHANAVADE")
    set_run_font(r, size=12, bold=True)
    r = body.add_run(
        " of Master in computer Application have successfully completed the Research Project work titled "
    )
    set_run_font(r, size=12)
    r = body.add_run(
        "‘RELAY: A MULTI-TENANT RETRIEVAL-AUGMENTED GENERATION PLATFORM FOR DOCUMENT-GROUNDED CHATBOTS’"
    )
    set_run_font(r, size=12, bold=True)
    r = body.add_run(
        " during the academic year 2026-27. This report is submitted as partial fulfillment of the requirement of degree in MCA Engineering of Modern College of Engineering."
    )
    set_run_font(r, size=12)
    anchor.addprevious(body._element)

    for _ in range(9):  # line space = 9 before signatures
        add_p("", size=12, line=1.0, after=0, before=0)

    add_p("External Examiner:", size=14, bold=True, align="left", before=6, after=18)

    sig = doc.add_table(rows=2, cols=3)
    names = ["Prof. Dr. Mrs. K. R. Joshi", "Prof. Dr. Shivani A. Budhkar", "Dr. Swati D. Ghule"]
    roles = ["Principal", "Head of Department", "Project Guide"]
    for j, (n, role) in enumerate(zip(names, roles)):
        c1 = sig.rows[0].cells[j]
        c2 = sig.rows[1].cells[j]
        c1.text = ""
        c2.text = ""
        p1 = c1.paragraphs[0]
        p2 = c2.paragraphs[0]
        fmt(p1, align="center", before=0, after=0, line=1.0)
        fmt(p2, align="center", before=2, after=0, line=1.0)
        r = p1.add_run(n)
        set_run_font(r, size=14, bold=True)
        r = p2.add_run(role)
        set_run_font(r, size=14, bold=True)
        clear_cell_borders(c1)
        clear_cell_borders(c2)
    anchor.addprevious(sig._tbl)
    add_bookmark(heading, "bm_certificate")


def remove_between(doc, start_prefix, end_prefix, keep_pred=None):
    start = find_para(doc, start_prefix)
    end = find_para(doc, end_prefix)
    for el in elements_until(start._element, lambda e: e is end._element):
        if el.tag == qn("w:p") and keep_pred and keep_pred(para_el_text(el)):
            continue
        delete_el(el)
    return end


def strip_text_diagrams(doc, paths):
    # 5.3 layers — insert figure, keep numbered explanation
    h = find_para(doc, "5.3 Proposed architecture")
    nxt = h._element.getnext()
    insert_figure_before(doc, nxt, paths["5.1"], "Figure 5.1: Proposed System Architecture")

    h = find_para(doc, "5.5 Document ingestion")
    insert_figure_before(doc, h._element.getnext(), paths["5.2"], "Figure 5.2: Document Ingestion Flowchart")

    h = find_para(doc, "5.6 Query methodology")
    insert_figure_before(doc, h._element.getnext(), paths["5.3"], "Figure 5.3: Query-Time RAG Flowchart")

    # 6.1 ASCII table sits between 6.1 heading and 6.2 heading
    end = remove_between(doc, "6.1 System architecture", "6.2 Use-case")
    insert_figure_before(doc, end._element, paths["6.1"], "Figure 6.1: System Architecture Diagram")

    end = remove_between(
        doc,
        "6.2 Use-case diagram",
        "6.3 Sequence",
        keep_pred=lambda t: t.startswith("Owner use cases include"),
    )
    insert_figure_before(doc, end._element, paths["6.2"], "Figure 6.2: Use Case Diagram")

    end = remove_between(doc, "6.3 Sequence", "6.4 Sequence")
    insert_figure_before(doc, end._element, paths["6.3"], "Figure 6.3: Sequence Diagram — Document Upload")

    end = remove_between(doc, "6.4 Sequence", "6.5 Sequence")
    insert_figure_before(doc, end._element, paths["6.4"], "Figure 6.4: Sequence Diagram — Dashboard Question and Answer")

    end = remove_between(doc, "6.5 Sequence", "6.6 Entity")
    insert_figure_before(doc, end._element, paths["6.5"], "Figure 6.5: Sequence Diagram — Public Widget Chat")

    end = remove_between(
        doc,
        "6.6 Entity-relationship",
        "6.7 Class diagram",
        keep_pred=lambda t: t.startswith("Pinecone is not a relational"),
    )
    insert_figure_before(doc, end._element, paths["6.6"], "Figure 6.6: Entity-Relationship Diagram")

    end = remove_between(doc, "6.7 Class diagram", "6.8 Activity")
    insert_figure_before(doc, end._element, paths["6.7"], "Figure 6.7: Class Diagram (Core Domain)")

    end = remove_between(doc, "6.8 Activity diagram", "6.9 Data-flow")
    insert_figure_before(doc, end._element, paths["6.8"], "Figure 6.8: Activity Diagram of RAG Answering")

    h = find_para(doc, "6.9 Data-flow")
    # insert after the explanation paragraph (the next paragraph), before chapter 7
    ch7 = find_para(doc, "7. PROJECT REQUIREMENT")
    insert_figure_before(doc, ch7._element, paths["6.9"], "Figure 6.9: Data-Flow Diagram — Ingestion and Query")


def set_roman_section(section):
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag == qn("w:pgNumType"):
            sectPr.remove(child)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), "lowerRoman")
    pg.set(qn("w:start"), "1")
    sectPr.append(pg)


def add_page_field(paragraph):
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el
    r1 = paragraph.add_run()
    r1._r.append(fld("begin"))
    r2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r2._r.append(instr)
    set_run_font(r2, size=12)
    r3 = paragraph.add_run()
    r3._r.append(fld("end"))


def setup_front_footer(doc):
    sec = doc.sections[0]
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    p.clear()
    fmt(p, align="right", before=0, after=0, line=1.0)
    add_page_field(p)
    set_roman_section(sec)


TOC = [
    ("Certificate", "bm_certificate"),
    ("Acknowledgement", "bm_ack"),
    ("Abstract", "bm_abs"),
    ("List of Abbreviations", "bm_abbr"),
    ("List of Figures", "bm_figs"),
    ("List of Tables", "bm_tabs"),
    ("1    Introduction, Aims, Motivation and Objectives", "bm_ch1"),
    ("2    Literature Survey", "bm_ch2"),
    ("3    Research Gap", "bm_ch3"),
    ("4    Problem Statement / Definition", "bm_ch4"),
    ("5    Proposed System / Proposed Methodology", "bm_ch5"),
    ("6    Diagrams (Architecture, Flowcharts, UML, ERD)", "bm_ch6"),
    ("7    Project Requirement Specification", "bm_ch7"),
    ("8    System Implementation and Code Documentation", "bm_ch8"),
    ("9    Result / Outcome / Experimental Result", "bm_ch9"),
    ("10  System Testing", "bm_ch10"),
    ("11  GUI / Screenshots", "bm_ch11"),
    ("12  Conclusions", "bm_ch12"),
    ("13  Limitations and Future Scope", "bm_ch13"),
    ("14  References", "bm_ch14"),
    ("15  Plagiarism Report", "bm_ch15"),
]

HEADINGS = [
    ("ACKNOWLEDGEMENT", "bm_ack"),
    ("ABSTRACT", "bm_abs"),
    ("LIST OF ABBREVIATIONS", "bm_abbr"),
    ("LIST OF FIGURES", "bm_figs"),
    ("LIST OF TABLES", "bm_tabs"),
    ("1. INTRODUCTION", "bm_ch1"),
    ("2. LITERATURE", "bm_ch2"),
    ("3. RESEARCH GAP", "bm_ch3"),
    ("4. PROBLEM STATEMENT", "bm_ch4"),
    ("5. PROPOSED SYSTEM", "bm_ch5"),
    ("6. FLOWCHART", "bm_ch6"),
    ("7. PROJECT REQUIREMENT", "bm_ch7"),
    ("8. SYSTEM IMPLEMENTATION", "bm_ch8"),
    ("9. RESULT", "bm_ch9"),
    ("10. SYSTEM TESTING", "bm_ch10"),
    ("11. GUI", "bm_ch11"),
    ("12. CONCLUSIONS", "bm_ch12"),
    ("13. LIMITATIONS", "bm_ch13"),
    ("14. REFERENCES", "bm_ch14"),
    ("15. PLAGIARISM", "bm_ch15"),
]


def add_pageref(paragraph, bookmark):
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el
    r1 = paragraph.add_run()
    r1._r.append(fld("begin"))
    r2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark} \\h "
    r2._r.append(instr)
    set_run_font(r2, size=12)
    r3 = paragraph.add_run()
    r3._r.append(fld("separate"))
    r4 = paragraph.add_run(" ")
    set_run_font(r4, size=12)
    r5 = paragraph.add_run()
    r5._r.append(fld("end"))


def rebuild_toc(doc):
    for prefix, bm in HEADINGS:
        add_bookmark(find_para(doc, prefix), bm)
    contents = find_para(doc, "CONTENTS")
    # remove the contents table (first table in body after CONTENTS)
    el = contents._element.getnext()
    while el is not None and el.tag != qn("w:tbl"):
        el = el.getnext()
    if el is None or el.tag != qn("w:tbl"):
        raise RuntimeError("Contents table not found")
    tbl = el
    # insert TOC entries before the table, then delete table
    for title, bm in TOC:
        p = doc.add_paragraph()
        fmt(p, align="left", before=2, after=2, line=1.5)
        pPr = p._element.get_or_add_pPr()
        tabs = pPr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            pPr.append(tabs)
        else:
            for child in list(tabs):
                tabs.remove(child)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "9020")
        tabs.append(tab)
        r = p.add_run(title)
        set_run_font(r, size=12)
        r = p.add_run("\t")
        set_run_font(r, size=12)
        add_pageref(p, bm)
        tbl.addprevious(p._element)
    delete_el(tbl)


def main():
    paths = build_figures()
    doc = Document(str(SRC))
    rebuild_certificate(doc)
    strip_text_diagrams(doc, paths)
    setup_front_footer(doc)
    rebuild_toc(doc)
    doc.save(str(DST))
    doc.save(str(COPY))
    print("saved", DST)


if __name__ == "__main__":
    main()
